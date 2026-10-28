#!/usr/bin/env python3
"""Cross-document traceability for one system: master.json integrity, per-document ID checks,
and the トレーサビリティマトリクス xlsx.

The design data lives next to the deliverables, always in <project dir>/<システム名>/ (the output root is
fixed to this repository). Every xlsx gets a _YYYYMMDDHHMMSS suffix, so each generation is a new file
(the matrix is the exception: one file, refreshed on every run).
Documents are written in this order, and generate.py refuses to skip ahead (missing_prereqs):
要件定義 → 基本設計 (per F) → 詳細仕様 (per F) → 詳細設計 (per SCR) → コード・テスト結果報告書 (per system)
    <出力ルート>/<システム名>/_design/master.json                  the single source of IDs and names
    <出力ルート>/<システム名>/_design/requirements-definition/requirements.json
    <出力ルート>/<システム名>/_design/basic-design/F-xx.json       one per function
    <出力ルート>/<システム名>/_design/detailed-spec/F-xx.json      one per function (tool spec)
    <出力ルート>/<システム名>/_design/detailed-spec/SCR-xx.json    screen-definition skeletons (optional)
    <出力ルート>/<システム名>/_design/detailed-design/SCR-xx.json  one per screen
    <出力ルート>/<システム名>/_design/code-generation/tests.json   test-case catalog (code-generation skill)

Usage:
    python traceability.py systems                       # output root + every system that has a master
    python traceability.py check-master <master.json>
    python traceability.py report <master.json> [--out matrix.xlsx]
    python traceability.py code-dir <master.json>        # CODE_DIR NEW|RESUME <dir> for the code-generation skill

Exit codes: 0 ok, 1 master.json is invalid.
"""
import argparse
import json
import re
import sys
from pathlib import Path

import consistency
import jsonschema
import test_results
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

REPO = Path(__file__).resolve().parent.parent
MASTER_SCHEMA = REPO / "templates" / "_shared" / "master.schema.json"

# Matches REQ-001 / F-01 / SCR-01 / IF-01 / TBL-01 / MSG-001 even when glued to Japanese text,
# but not the "F-01" inside "IF-01".
ID_RE = re.compile(r"(?<![A-Za-z0-9])(?:REQ-\d{3}|F-\d{2}|SCR-\d{2}|IF-\d{2}|TBL-\d{2}|MSG-\d{3})(?![0-9])")

KINDS = [  # (master key, label)
    ("requirements", "要件"), ("functions", "機能"), ("screens", "画面"),
    ("interfaces", "IF"), ("tables", "テーブル"), ("messages", "メッセージ"),
]

# _design/<dir> -> (output folder under <出力ルート>/<システム名>/, xlsx name suffixes, label), in writing order
DOCS = {
    "requirements-definition": ("02要件定義", [""], "要件定義書"),
    "basic-design": ("03基本設計", [""], "基本設計書"),
    "detailed-spec": ("04詳細仕様", ["_ツール仕様書"], "詳細仕様書"),
    "detailed-design": ("05詳細設計", ["_IF定義書", "_詳細設計書"], "詳細設計書"),
    "code-generation": ("06テスト結果", ["_テスト結果報告書"], "テスト結果報告書"),
}
SYSTEM_DOCS = {"requirements-definition": "requirements", "code-generation": "tests"}  # one per system


def doc_root():
    """The output root: always the project directory (this repository)."""
    return REPO


def design_dir(system_name):
    return doc_root() / system_name / "_design"


CODE_DIR_PREFIX = "07ソースコード_"  # <project dir>/<システム名>/07ソースコード_<timestamp>/ (code-generation skill)


def code_dir(master):
    """Where the code-generation skill writes code: the newest 07ソースコード_* folder of the system if its
    .design-codegen.json still lacks some F (resume), else a new folder with a fresh timestamp."""
    from datetime import datetime
    sys_root = doc_root() / master["system"]["name"]
    all_f = {f["id"] for f in master["functions"]}
    for d in sorted(sys_root.glob(f"{CODE_DIR_PREFIX}*"), reverse=True):
        state = d / ".design-codegen.json"
        if state.exists():
            done = set(json.loads(state.read_text(encoding="utf-8")).get("implemented", []))
            if not all_f <= done:
                return "RESUME", d
        break
    return "NEW", sys_root / f"{CODE_DIR_PREFIX}{datetime.now().strftime('%Y%m%d%H%M%S')}"


def load_master(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def all_ids(master):
    return {item["id"]: item for key, _ in KINDS for item in master[key]}


def check_master(master):
    """Schema + referential integrity. Returns a list of error strings."""
    schema = json.loads(MASTER_SCHEMA.read_text(encoding="utf-8"))
    errors = [f"{'/'.join(map(str, e.absolute_path)) or '(root)'}: {e.message}"
              for e in jsonschema.Draft7Validator(schema).iter_errors(master)]
    if errors:
        return errors

    seen = {}
    for key, _ in KINDS:
        for item in master[key]:
            if item["id"] in seen:
                errors.append(f"{item['id']}: ID が重複している")
            seen[item["id"]] = item

    def expect(owner, ref, prefix):
        if not ref.startswith(prefix) or ref not in seen:
            errors.append(f"{owner}: 参照先 {ref} が主表に存在しない（{prefix}xx を期待）")

    for f in master["functions"]:
        for r in f["requirements"]:
            expect(f["id"], r, "REQ-")
    for s in master["screens"]:
        for r in s["functions"]:
            expect(s["id"], r, "F-")
    for i in master["interfaces"]:
        if i["screen"] is not None:
            expect(i["id"], i["screen"], "SCR-")
        for r in i["functions"]:
            expect(i["id"], r, "F-")
    for key in ("tables", "messages"):
        for t in master[key]:
            for r in t["functions"]:
                expect(t["id"], r, "F-")

    for d in master["initial_data"]:
        expect("initial_data", d["table"], "TBL-")
    if any(i["auth"] for i in master["interfaces"]) and not master["initial_data"]:
        errors.append("initial_data: 認証が必要な IF があるのに、ログインに使う初期ユーザーが定義されていない")

    covered = {r for f in master["functions"] for r in f["requirements"]}
    for req in master["requirements"]:
        if req["id"] not in covered:
            errors.append(f"{req['id']}: どの機能(F)からも参照されていない（要件が設計に落ちていない）")
    return errors


def expected_refs(master, doc_dir, target):
    """IDs a document must mention, by the cross-document rules in master_rules.md."""
    if doc_dir == "requirements-definition":
        return [r["id"] for r in master["requirements"]]
    if doc_dir == "code-generation":
        return [x["id"] for key in ("functions", "screens", "interfaces", "messages") for x in master[key]]
    if target.startswith("F-"):
        func = next(f for f in master["functions"] if f["id"] == target)
        if doc_dir == "basic-design":
            return [target, *func["requirements"]]
        # detailed-spec tool spec: its screens and its messages
        return [target,
                *[s["id"] for s in master["screens"] if target in s["functions"]],
                *[m["id"] for m in master["messages"] if target in m["functions"]]]
    if target.startswith("SCR-") and doc_dir == "detailed-design":
        screen = next(s for s in master["screens"] if s["id"] == target)
        return [target,
                *[i["id"] for i in master["interfaces"] if i["screen"] == target],
                *[m["id"] for m in master["messages"] if set(m["functions"]) & set(screen["functions"])]]
    return [target]


def check_content(master, doc_dir, target, content):
    """Errors for one document's content (without _notes_for_review). Empty list = consistent."""
    ids = all_ids(master)
    errors = []
    if doc_dir not in SYSTEM_DOCS and target not in ids:
        return [f"対象 {target} が主表に存在しない（ファイル名は主表の ID にする）"]
    text = json.dumps(content, ensure_ascii=False)
    used = set(ID_RE.findall(text))
    for unknown in sorted(used - ids.keys()):
        errors.append(f"{unknown} が主表に存在しない（先に主表へ追記するか、既存の ID を使う）")
    for missing in expected_refs(master, doc_dir, target):
        if missing not in used:
            errors.append(f"{missing}（{ids[missing].get('name') or ids[missing].get('text')}）への言及が無い")
    sys_name = master["system"]["name"]
    if content.get("meta", {}).get("system_name") != sys_name:
        errors.append(f"meta.system_name が主表のシステム名「{sys_name}」と一致しない")
    return errors


def generated(master, doc_dir, target):
    """Whether the xlsx for (doc_dir, target) exists in the default output folder (any timestamp)."""
    folder, suffixes, _ = DOCS[doc_dir]
    out = doc_root() / master["system"]["name"] / folder
    if doc_dir in SYSTEM_DOCS:
        return all(any(out.glob(f"{master['system']['name']}{suffix}_*.xlsx")) for suffix in suffixes)
    return all(any(out.glob(f"{target}_*{suffix}_*.xlsx")) for suffix in suffixes)


def doc_targets(master):
    return {
        "requirements-definition": ["requirements"],
        "basic-design": [f["id"] for f in master["functions"]],
        "detailed-spec": [f["id"] for f in master["functions"]],
        "detailed-design": [s["id"] for s in master["screens"]],
        "code-generation": ["tests"],
    }


def missing_prereqs(master, sys_dir, doc_dir, target):
    """Upstream documents that must already exist before (doc_dir, target) is written:
    基本設計 F needs the 要件定義, 詳細仕様 F needs the 基本設計 of the same F, and 詳細設計 SCR needs
    the 詳細仕様 of every F on that screen. Returns their labels (empty when ready)."""
    if doc_dir == "basic-design":
        needed = [("requirements-definition", "requirements")]
    elif doc_dir == "detailed-spec" and target.startswith("F-"):
        needed = [("basic-design", target)]
    elif doc_dir == "code-generation":
        needed = [*(("detailed-spec", f["id"]) for f in master["functions"]),
                  *(("detailed-design", s["id"]) for s in master["screens"])]
    elif doc_dir == "detailed-design":
        screen = next((s for s in master["screens"] if s["id"] == target), None)
        needed = [("detailed-spec", f) for f in (screen["functions"] if screen else [])]
    else:
        needed = []
    return [f"{DOCS[d][2]} {t}" for d, t in needed if not (sys_dir / d / f"{t}.json").exists()]


def collect_issues(master, sys_dir):
    issues = [("主表", "master.json", e) for e in check_master(master)]
    if issues:
        return issues
    for doc_dir, wanted in doc_targets(master).items():
        label = DOCS[doc_dir][2]
        for target in wanted:
            path = sys_dir / doc_dir / f"{target}.json"
            if not path.exists():
                issues.append(("未作成", f"{label} {target}", "まだ作成されていない"))
                continue
            content = json.loads(path.read_text(encoding="utf-8"))
            content.pop("_notes_for_review", None)
            errors = check_content(master, doc_dir, target, content) or consistency.check(master, sys_dir, doc_dir, target, content)
            for e in errors:
                issues.append(("不整合", f"{label} {target}", e))
    return issues


def write_matrix(master, sys_dir, out_path):
    issues = collect_issues(master, sys_dir)
    thin = Side(style="thin", color="808080")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="DDEBF7")
    font = Font(name="ＭＳ Ｐゴシック", size=10)
    wrap = Alignment(wrap_text=True, vertical="top")

    def sheet(ws, headers, rows, widths):
        ws.append(headers)
        for r in rows:
            ws.append(r)
        for row in ws.iter_rows():
            for c in row:
                c.font, c.border, c.alignment = font, border, wrap
        for c in ws[1]:
            c.fill, c.font = head_fill, Font(name="ＭＳ Ｐゴシック", size=10, bold=True)
        for i, w in enumerate(widths):
            ws.column_dimensions[chr(ord("A") + i)].width = w
        ws.freeze_panes = "A2"

    def done(doc_dir, target):
        return "済" if generated(master, doc_dir, target) else "未"

    tests_path = sys_dir / "code-generation" / "tests.json"
    test_summary = {}
    if tests_path.exists() and generated(master, "code-generation", "tests"):
        rows, _ = test_results.results(json.loads(tests_path.read_text(encoding="utf-8")))
        test_summary = test_results.summary_by_function(rows)

    def tested(func_id):
        s = test_summary.get(func_id)
        if not s:
            return "未"
        return f"OK {s['OK']}/{s['total']}" + (f"（NG {s['NG']}）" if s["NG"] else "") + (f"（未実施 {s['未実施']}）" if s["未実施"] else "")

    rows = []
    for req in master["requirements"]:
        funcs = [f for f in master["functions"] if req["id"] in f["requirements"]] or [None]
        for f in funcs:
            if f is None:
                rows.append([req["id"], req["name"], req["category"], "", "", "", "", "", "", "", "", ""])
                continue
            screens = [s for s in master["screens"] if f["id"] in s["functions"]]
            ifs = [i for i in master["interfaces"] if f["id"] in i["functions"]]
            rows.append([
                req["id"], req["name"], req["category"], f["id"], f["name"],
                done("basic-design", f["id"]), done("detailed-spec", f["id"]),
                "\n".join(f"{s['id']} {s['name']}" for s in screens),
                "\n".join(f"{s['id']}：{done('detailed-design', s['id'])}" for s in screens),
                "\n".join(f"{i['id']} {i['method']} {i['path']}" for i in ifs),
                tested(f["id"]),
                done("requirements-definition", "requirements"),
            ])

    wb = Workbook()
    ws = wb.active
    ws.title = "トレーサビリティ"
    sheet(ws, ["要件ID", "要件名", "分類", "機能ID", "機能名", "基本設計書", "詳細仕様書",
               "画面", "詳細設計書", "IF", "テスト結果", "要件定義書"], rows,
          [10, 28, 12, 8, 22, 10, 10, 22, 14, 34, 16, 10])

    sheet(wb.create_sheet("不整合・未作成"), ["種別", "対象", "内容"],
          issues or [("", "", "不整合・未作成なし")], [10, 26, 80])

    id_rows = []
    for key, label in KINDS:
        for item in master[key]:
            name = item.get("name") or item.get("text")
            refs = item.get("requirements") or item.get("functions") or []
            id_rows.append([label, item["id"], name, ", ".join(refs)])
    sheet(wb.create_sheet("ID一覧"), ["種類", "ID", "名称", "関連"], id_rows, [12, 10, 40, 30])

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return issues


def default_matrix_path(master):
    name = master["system"]["name"]
    return doc_root() / name / f"{name}_トレーサビリティマトリクス.xlsx"


def summarize(master):
    for key, label in KINDS:
        items = master[key]
        if items:
            print(f"{label}（{len(items)}件）:")
            for item in items:
                refs = item.get("requirements") or item.get("functions") or []
                print(f"  {item['id']} {item.get('name') or item.get('text')}" + (f"  ← {', '.join(refs)}" if refs else ""))


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("command", choices=["systems", "check-master", "report", "code-dir"])
    parser.add_argument("master", type=Path, nargs="?")
    parser.add_argument("--out", type=Path, help="matrix xlsx path (report only)")
    args = parser.parse_args()

    if args.command == "systems":
        root = doc_root()
        print(f"DOC_ROOT {root}")
        for master_path in sorted(root.glob("*/_design/master.json")):
            print(f"SYSTEM {master_path.parent.parent.name}\t{master_path}")
        print(f"NEW_SYSTEM_MASTER {design_dir('<システム名>') / 'master.json'}")
        return 0
    if args.master is None:
        parser.error(f"{args.command} needs the master.json path")

    master = load_master(args.master)
    errors = check_master(master)
    if errors:
        print("MASTER INVALID:", file=sys.stderr)
        for e in errors:
            print(f"- {e}", file=sys.stderr)
        return 1

    if args.command == "code-dir":
        mode, path = code_dir(master)
        print(f"CODE_DIR {mode} {path}")
        return 0
    if args.command == "check-master":
        print("MASTER OK")
        summarize(master)
        return 0

    out = args.out or default_matrix_path(master)
    issues = write_matrix(master, args.master.parent, out)
    print(f"WROTE {out}")
    for kind in ("不整合", "未作成"):
        hits = [i for i in issues if i[0] == kind]
        if hits:
            print(f"{kind}（{len(hits)}件）:")
            for _, target, msg in hits:
                print(f"- {target}: {msg}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
