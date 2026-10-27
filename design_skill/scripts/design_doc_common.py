"""Shared rendering primitives for every 日本語SI設計書 category (基本設計,
要件定義, 詳細設計, 詳細仕様...). All of them share the same house style
(方眼紙 grid, MS Gothic 10pt, thin-bordered tables, a page header row with
シート#/システム名/機能名/作成者/作成日/改訂者/改訂日/版) captured once from
templates/_shared/house_style.xlsx (a copy of the original 基本設計 sample).

A category-specific renderer script:
1. calls load_house_style() to get a fresh workbook + captured style dict
2. calls prepare_chapter_sheets() to reshape the six template chapter sheets
   into however many chapters that category actually needs
3. writes chapter content using the emit_* primitives below (adding new
   primitives in the category's own script when a category needs a shape
   basic-design didn't have, e.g. detailed-design's step blocks)
4. calls render_toc / render_revision_history / sync_page_header as needed
5. saves with save_workbook(), never wb.save() — openpyxl drops the 表紙 sheet's drawing
   (title text box and 承認/発行 stamp frame), and save_workbook() puts it back
"""
import copy
import re
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

HOUSE_STYLE_PATH = Path(__file__).resolve().parent.parent / "templates" / "_shared" / "house_style.xlsx"
SOURCE_CHAPTER_SHEETS = [
    "1.本書の定義", "2.機能の概要", "3.ログ出力機構要件",
    "4.ログ出力機構の機能", "5.機能詳細", "6.ログ仕様",
]

HEADER_COORDS = {
    "system_name": "O2",
    "function_name": "T2",
    "author": "AS2",
    "created_date": "AX2",
    "reviser": "BC2",
    "revised_date": "BH2",
    "version": "BM2",
}
TITLE_COORD = "C1"

INVALID_SHEETNAME_CHARS = re.compile(r"[\\/*?:\[\]]")


def sanitize_sheet_title(text: str) -> str:
    return INVALID_SHEETNAME_CHARS.sub("", text)[:31]


def style_of(cell):
    return dict(
        font=copy.copy(cell.font),
        alignment=copy.copy(cell.alignment),
        fill=copy.copy(cell.fill),
        border=copy.copy(cell.border),
        number_format=cell.number_format,
    )


def apply_style(cell, style):
    cell.font = style["font"]
    cell.alignment = style["alignment"]
    cell.fill = style["fill"]
    cell.border = style["border"]
    cell.number_format = style["number_format"]


def capture_styles(wb):
    """Read style snapshots from the *original* house-style sheet names.
    Must be called before prepare_chapter_sheets() renames/clones/deletes them.
    """
    ch1 = wb["1.本書の定義"]
    ch5 = wb["5.機能詳細"]
    ch6 = wb["6.ログ仕様"]
    toc = wb["目次"]
    return {
        "chapter_b": style_of(ch1["B5"]),
        "chapter_c": style_of(ch1["C5"]),
        "subsection_num": style_of(ch1["C6"]),
        "subsection_title": style_of(ch1["D6"]),
        "body_line": style_of(ch1["D9"]),
        "item_title": style_of(ch1["D37"]),
        "item_body": style_of(ch1["E38"]),
        "table_header": style_of(ch5["D25"]),
        "table_data": style_of(ch5["D26"]),
        "table3_header": style_of(ch6["D13"]),
        "table3_data": style_of(ch6["D14"]),
        "toc_chapter_num": style_of(toc["C7"]),
        "toc_chapter_title": style_of(toc["D7"]),
        "toc_sub_num": style_of(toc["D8"]),
        "toc_sub_title": style_of(toc["E8"]),
        "section_mark": style_of(toc["B5"]),
        "section_title": style_of(toc["C5"]),
    }


def load_house_style():
    wb = openpyxl.load_workbook(HOUSE_STYLE_PATH)
    styles = capture_styles(wb)
    return wb, styles


def clear_body(ws):
    """Delete everything from row 3 downward; rows 1-2 (page header) survive.

    openpyxl's delete_rows does not clean up merged-cell ranges that fall
    entirely within the deleted rows, so body-area merges must be unmerged
    first, or writes to their non-anchor cells silently land on read-only
    MergedCell objects.
    """
    for rng in [r for r in list(ws.merged_cells.ranges) if r.min_row >= 3]:
        ws.unmerge_cells(str(rng))
    if ws.max_row >= 3:
        ws.delete_rows(3, ws.max_row - 2)


def prepare_chapter_sheets(wb, chapter_sheet_names, include_toc=True):
    """Reshape the workbook's template chapter sheets into exactly
    len(chapter_sheet_names) sheets, in order, cleared and ready to render
    into. Returns the worksheets in order. Drops 目次 if include_toc=False.
    """
    if not include_toc and "目次" in wb.sheetnames:
        del wb["目次"]

    existing = [n for n in SOURCE_CHAPTER_SHEETS if n in wb.sheetnames]
    needed = len(chapter_sheet_names)

    for name in existing[needed:]:
        del wb[name]
    existing = existing[:needed]

    donor_name = existing[-1] if existing else None
    while len(existing) < needed:
        donor = wb[donor_name]
        new_ws = wb.copy_worksheet(donor)
        existing.append(new_ws.title)
        donor_name = new_ws.title

    sheets = []
    for old_name, new_name in zip(existing, chapter_sheet_names):
        ws = wb[old_name]
        clear_body(ws)
        ws.title = sanitize_sheet_title(new_name)
        sheets.append(ws)
    return sheets


def set_cell(ws, row, col_letter, value, style):
    cell = ws.cell(row=row, column=column_index_from_string(col_letter), value=value)
    apply_style(cell, style)
    return cell


def ensure_bullet(text: str) -> str:
    text = text.strip()
    return text if text.startswith("・") else f"・{text}"


def ensure_prefix(text: str, marker: str) -> str:
    text = text.strip()
    return text if text.startswith(marker) else f"{marker}{text}"


# ---------------------------------------------------------------------------
# Emission primitives
# ---------------------------------------------------------------------------

def emit_chapter_heading(ws, row, title, styles):
    set_cell(ws, row, "B", "1.", styles["chapter_b"])
    set_cell(ws, row, "C", title, styles["chapter_c"])
    return row + 1


def emit_numbered_subsection(ws, row, num, title, body_lines, styles):
    set_cell(ws, row, "C", num, styles["subsection_num"])
    set_cell(ws, row, "D", title, styles["subsection_title"])
    row += 1
    for line in body_lines:
        set_cell(ws, row, "D", line, styles["body_line"])
        row += 1
    return row + 1  # trailing blank spacer


def emit_item_group(ws, row, items, styles, bullet=False):
    for item in items:
        title = ensure_bullet(item["title"]) if bullet else item["title"]
        set_cell(ws, row, "D", title, styles["item_title"])
        row += 1
        for line in item["body"]:
            set_cell(ws, row, "E", line, styles["item_body"])
            row += 1
        row += 1  # spacer between items
    return row


def emit_table2(ws, row, col_headers, rows, styles, col_a="D", col_b="H"):
    set_cell(ws, row, col_a, col_headers[0], styles["table_header"])
    set_cell(ws, row, col_b, col_headers[1], styles["table_header"])
    row += 1
    for r in rows:
        set_cell(ws, row, col_a, r[0], styles["table_data"])
        set_cell(ws, row, col_b, r[1], styles["table_data"])
        row += 1
    return row + 1


def spaced_columns(n, start="D", step=3):
    """n column letters spaced `step` apart starting at `start` (default
    D, G, J, M...) — a reasonable default width for a fresh N-column table
    when the category doesn't need specific historical coordinates."""
    start_idx = column_index_from_string(start)
    return [get_column_letter(start_idx + i * step) for i in range(n)]


def emit_table_n(ws, row, headers, rows, styles, cols=None):
    """N-column table. `cols` is a list of column letters, one per header;
    defaults to spaced_columns(len(headers)) if not given."""
    cols = cols or spaced_columns(len(headers))
    for c, h in zip(cols, headers):
        set_cell(ws, row, c, h, styles["table3_header"])
    row += 1
    for r in rows:
        for c, v in zip(cols, r):
            set_cell(ws, row, c, v, styles["table3_data"])
        row += 1
    return row + 1


def emit_section_mark(ws, row, title, styles):
    """A '■<title>' level-1 heading, used by per-screen docs (IF定義,
    詳細設計書本体) that use ■/○ markers instead of basic-design's numbered
    '1.'/'1.1' scheme. Reuses the section_mark/section_title styles captured
    from 目次/改定履歴's own '■目次' heading."""
    set_cell(ws, row, "B", "■", styles["section_mark"])
    set_cell(ws, row, "C", title, styles["section_title"])
    return row + 1


def emit_sub_mark(ws, row, title, body_lines, styles):
    """A '○<title>' level-2 heading with plain body lines beneath it."""
    set_cell(ws, row, "D", ensure_prefix(title, "○"), styles["item_title"])
    row += 1
    for line in body_lines:
        set_cell(ws, row, "E", line, styles["item_body"])
        row += 1
    return row + 1


def emit_placeholder_note(ws, row, text, styles):
    """A single italic-less plain line flagging manual-fill content (e.g. a
    screenshot or hand-drawn diagram the renderer can't produce)."""
    set_cell(ws, row, "D", f"（{text}）", styles["body_line"])
    return row + 1


def render_toc(ws, chapters, styles):
    clear_body(ws)
    set_cell(ws, 5, "B", "■", styles["section_mark"])
    set_cell(ws, 5, "C", "目次", styles["section_title"])
    row = 7
    for i, (chapter_title, subsections) in enumerate(chapters, start=1):
        set_cell(ws, row, "C", f"{i}.", styles["toc_chapter_num"])
        set_cell(ws, row, "D", chapter_title, styles["toc_chapter_title"])
        row += 1
        for sub_num, sub_title in subsections:
            set_cell(ws, row, "D", sub_num, styles["toc_sub_num"])
            set_cell(ws, row, "E", sub_title, styles["toc_sub_title"])
            row += 1
        row += 1  # spacer between chapters


def render_revision_history(ws, meta, styles):
    """Fill the first row of the house-style 改定履歴 table in place. The table itself (header
    row 7, empty merged rows down to row 28 for later revisions) is kept exactly as in the template."""
    ws["C8"] = meta["version"]
    ws["E8"] = meta["created_date"]
    ws["K8"] = meta["author"]
    ws["Q8"] = "ー"
    ws["AD8"] = "新規作成"


def sync_page_header(wb, meta, sheet_names, title_label):
    """title_label is the full bracketed doc-type text, e.g. '要件定義書',
    'I/F定義書', 'ツール仕様書'. Every sheet repeats '[<title_label>]\\n改定履歴'
    in C1 — a quirk of the original house-style sample, kept for consistency.
    """
    title = f"[{title_label}]\n改定履歴"
    for name in sheet_names:
        ws = wb[name]
        ws[TITLE_COORD] = title
        ws[HEADER_COORDS["system_name"]] = meta["system_name"]
        ws[HEADER_COORDS["function_name"]] = meta["function_name"]
        ws[HEADER_COORDS["author"]] = meta["author"]
        ws[HEADER_COORDS["created_date"]] = meta["created_date"]
        ws[HEADER_COORDS["reviser"]] = meta["author"]
        ws[HEADER_COORDS["revised_date"]] = meta["created_date"]
        ws[HEADER_COORDS["version"]] = meta["version"]


# ---------------------------------------------------------------------------
# Saving (表紙 drawing)
# ---------------------------------------------------------------------------

COVER_SHEET = "表紙"
DRAWING_TYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing"
DRAWING_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.drawing+xml"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
COVER_TITLE_SHAPE = re.compile(r'(<xdr:sp>(?:(?!</xdr:sp>).)*?name="四角形 2".*?)<xdr:txBody>.*?</xdr:txBody>', re.S)
# Elements that must come after <drawing> in a worksheet (CT_Worksheet sequence).
AFTER_DRAWING = ("<legacyDrawing", "<legacyDrawingHF", "<drawingHF", "<picture", "<oleObjects",
                 "<controls", "<webPublishItems", "<tableParts", "<extLst", "</worksheet>")


def cover_lines(meta, doc_kind):
    """Title lines for the 表紙 text box: system name, document target (when it is not the
    system itself, e.g. 'F-01 ログイン機能'), and the document kind."""
    lines = [meta["system_name"]]
    if meta["function_name"] != meta["system_name"]:
        lines.append(meta["function_name"].replace("_", " ", 1))
    lines.append(doc_kind)
    return lines


def _cover_font_size(lines):
    """Largest size (pt, at most the template's 48) that fits the title box (≈488pt x 200pt)."""
    widest = max(sum(0.55 if ord(ch) < 0x2E80 else 1.0 for ch in line) for line in lines)
    return int(min(48, 480 / max(widest, 1), 200 / (1.3 * len(lines))))


def _cover_txbody(lines):
    size = _cover_font_size(lines) * 100
    run_pr = (f'lang="ja-JP" altLang="en-US" sz="{size}"><a:latin typeface="+mj-ea"/>'
              f'<a:ea typeface="+mj-ea"/>')
    paras = "".join(
        f'<a:p><a:pPr algn="ctr"/><a:r><a:rPr {run_pr}</a:rPr><a:t>{escape(line)}</a:t></a:r></a:p>'
        for line in lines)
    return ('<xdr:txBody><a:bodyPr vertOverflow="clip" horzOverflow="clip" wrap="square" rtlCol="0" '
            f'anchor="ctr" anchorCtr="0"/><a:lstStyle/>{paras}</xdr:txBody>')


def _cover_sheet_part(files):
    """Zip path of the 表紙 worksheet, resolved through workbook.xml and its rels."""
    workbook = files["xl/workbook.xml"].decode("utf-8")
    m = re.search(rf'<sheet [^>]*name="{COVER_SHEET}"[^>]*r:id="([^"]+)"', workbook)
    if not m:
        return None
    rels = files["xl/_rels/workbook.xml.rels"].decode("utf-8")
    rel = re.search(rf'<Relationship [^>]*Id="{m.group(1)}"[^>]*/>', rels).group(0)
    target = re.search(r'Target="([^"]+)"', rel).group(1)
    return target.lstrip("/") if target.startswith("/") else f"xl/{target}"


def save_workbook(wb, output_path, cover):
    """wb.save(), then restore the house-style 表紙 drawing with `cover` (lines from cover_lines())
    as the title. openpyxl cannot keep shapes, so the drawing part is copied from house_style.xlsx
    and wired back into the saved package (sheet <drawing>, sheet rels, content type)."""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)

    with zipfile.ZipFile(HOUSE_STYLE_PATH) as src:
        drawing = src.read("xl/drawings/drawing1.xml").decode("utf-8")
    drawing, n = COVER_TITLE_SHAPE.subn(lambda m: m.group(1) + _cover_txbody(cover), drawing, count=1)
    if n != 1:
        raise RuntimeError("house_style.xlsx: 表紙 title shape '四角形 2' not found")

    with zipfile.ZipFile(output_path) as z:
        files = {i.filename: z.read(i.filename) for i in z.infolist()}
    sheet_part = _cover_sheet_part(files)
    if sheet_part is None:
        return
    sheet_dir, sheet_file = sheet_part.rsplit("/", 1)
    rels_part = f"{sheet_dir}/_rels/{sheet_file}.rels"
    drawing_part = "xl/drawings/drawing_cover.xml"

    rels = files.get(rels_part, b"").decode("utf-8") or (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"></Relationships>')
    rels = rels.replace("</Relationships>",
                        f'<Relationship Id="rIdCover" Type="{DRAWING_TYPE}" Target="../drawings/drawing_cover.xml"/>'
                        "</Relationships>")

    sheet = files[sheet_part].decode("utf-8")
    if "xmlns:r=" not in sheet[:sheet.index(">")]:
        sheet = sheet.replace("<worksheet ", f'<worksheet xmlns:r="{R_NS}" ', 1)
    at = min(i for i in (sheet.find(tag) for tag in AFTER_DRAWING) if i >= 0)
    sheet = sheet[:at] + '<drawing r:id="rIdCover"/>' + sheet[at:]

    types = files["[Content_Types].xml"].decode("utf-8").replace(
        "</Types>", f'<Override PartName="/{drawing_part}" ContentType="{DRAWING_CONTENT_TYPE}"/></Types>')

    files.update({sheet_part: sheet.encode("utf-8"), rels_part: rels.encode("utf-8"),
                  drawing_part: drawing.encode("utf-8"), "[Content_Types].xml": types.encode("utf-8")})
    with zipfile.ZipFile(output_path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in files.items():
            z.writestr(name, data)
